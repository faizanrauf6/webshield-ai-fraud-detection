#!/usr/bin/env python3
"""Build independent final-report/manual exports without touching proposal files.

Dependencies: python-docx, reportlab, Markdown, Pillow (see official package docs).
Run from any directory: python scripts/build-submission.py
"""
from __future__ import annotations

import base64
import hashlib
import html
import json
import re
from pathlib import Path

import markdown
import reportlab
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import html as html_parser
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus.tableofcontents import TableOfContents
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, PageBreak, Image as PDFImage,
    Table, TableStyle, KeepTogether, Preformatted,
)

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
OUT = DOCS / 'submission'
ASSETS = DOCS / 'assets'
META = json.loads((OUT / 'sample-cover-metadata.json').read_text())
FONT_DIR = Path('/System/Library/Fonts/Supplemental')


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def diagram_assets():
    """Draw readable implementation diagrams with explicit trust boundaries."""
    regular = str(FONT_DIR / 'Arial.ttf')
    bold = str(FONT_DIR / 'Arial Bold.ttf')
    if not Path(regular).exists():
        bundled_fonts = Path(reportlab.__file__).parent / 'fonts'
        regular = str(bundled_fonts / 'Vera.ttf')
        bold = str(bundled_fonts / 'VeraBd.ttf')
    def font(size, heavy=False):
        return ImageFont.truetype(bold if heavy else regular, size)
    def canvas(title, subtitle, height):
        im = Image.new('RGB', (1500, height), '#f7fafb')
        d = ImageDraw.Draw(im)
        d.text((60,35), title, fill='#16344a', font=font(38, True))
        d.text((60,90), subtitle, fill='#445b6d', font=font(22))
        return im,d
    def box(d, xy, title, lines, color='#e7f4ef'):
        d.rounded_rectangle(xy, 16, fill=color, outline='#467160', width=3)
        x,y,x2,y2=xy
        heading_size = 24
        while d.textlength(title, font=font(heading_size, True)) > x2-x-40 and heading_size > 17:
            heading_size -= 1
        d.text((x+20,y+17),title,fill='#16344a',font=font(heading_size,True))
        for i,line in enumerate(lines):
            d.text((x+20,y+55+i*29),line,fill='#334c5e',font=font(21))
    def arrow(d,a,b):
        d.line([a,b],fill='#587184',width=5)
        import math
        angle=math.atan2(b[1]-a[1],b[0]-a[0])
        pts=[b,(b[0]-17*math.cos(angle-.5),b[1]-17*math.sin(angle-.5)),(b[0]-17*math.cos(angle+.5),b[1]-17*math.sin(angle+.5))]
        d.polygon(pts,fill='#587184')
    im,d=canvas('WebShield AI architecture','Local application · isolated static collector · offline model development',880)
    box(d,(60,170,355,290),'Browser / Next.js',['Same-origin API proxy','Reports and history'])
    box(d,(470,170,800,330),'NestJS API + queue worker',['Validate → persist → enqueue','Bounded orchestration'])
    arrow(d,(355,230),(470,230))
    box(d,(920,160,1440,280),'PostgreSQL 16',['Saved lifecycle and JSONB evidence'])
    box(d,(920,320,1440,440),'Redis 7 / BullMQ',['Persistent jobs and retries'])
    arrow(d,(800,215),(920,215));arrow(d,(800,290),(920,360))
    box(d,(60,430,410,590),'Trusted Python inference',['Installed lexical-v2 bundle','30-second hard deadline'],'#eaf0fb')
    arrow(d,(580,330),(410,465))
    box(d,(500,500,990,690),'Isolated Python container',['Kernel egress firewall; unprivileged','Public DNS / pinned static HTML','Authoritative RDAP; no scripts'],'#fff1df')
    arrow(d,(640,330),(640,500))
    box(d,(1050,540,1440,690),'Fixed provider endpoints',['Google / OpenPhish optional','Exact URL + time + expiry'],'#eaf0fb')
    d.line([(800,310),(865,310),(865,500)],fill='#587184',width=5)
    arrow(d,(865,500),(1120,540))
    box(d,(60,690,410,810),'Offline research',['Import → train → evaluate','Trusted-local promotion'],'#eaf0fb')
    arrow(d,(230,690),(230,590))
    d.text((500,770),'Target pages are never executed in the user browser.',fill='#445b6d',font=font(23))
    d.text((500,805),'Collection and provider findings remain attributable.',fill='#445b6d',font=font(23))
    im.save(ASSETS/'architecture.png')
    im,d=canvas('WebShield AI database design','PostgreSQL persistence · queue jobs live separately in Redis',880)
    box(d,(60,165,850,720),'scans',[
        'PK  id : uuid',
        'url_input, url_normalized, hostname : text',
        'status : queued | running | completed | partial | failed',
        'fixture : text, optional synthetic scenario',
        'analysis : jsonb, assessment + evidence + model version',
        'error_code : text, optional',
        'created_at, started_at, finished_at : timestamptz(3)',
        '',
        'Indexes: (created_at DESC, id DESC), status, hostname',
        'Saved before enqueue; historical observations preserved',
    ])
    box(d,(925,165,1440,405),'model_versions',[
        'PK  version : text',
        'metadata : jsonb',
        'created_at : timestamptz',
        '',
        'Version metadata catalog',
    ],'#eaf0fb')
    box(d,(925,470,1440,720),'Redis / BullMQ',[
        'Queue jobs and retry state',
        'References scan UUID',
        'Startup reconciles with PostgreSQL',
        '',
        'Separate persistence service',
    ],'#fff1df')
    d.text((60,770),'No foreign key links scans to model_versions.',fill='#445b6d',font=font(24,True))
    d.text((60,810),'Assessment JSON records the model version; trusted local bundle controls runtime loading.',fill='#445b6d',font=font(22))
    im.save(ASSETS/'database.png')


def source_tree(path):
    markup=markdown.markdown(path.read_text(),extensions=['tables','fenced_code','toc'])
    return html_parser.fragment_fromstring(markup,create_parent='article')


def text_of(el):
    return ''.join(el.itertext()).strip()


def inline(el):
    """Convert safe Markdown inline elements to ReportLab paragraph markup."""
    s=html.escape(el.text or '')
    for child in el:
        inner=inline(child)
        if child.tag in ('strong','b'): inner=f'<b>{inner}</b>'
        elif child.tag in ('em','i'): inner=f'<i>{inner}</i>'
        elif child.tag=='code': inner=f'<font name="Courier" size="8.5">{inner}</font>'
        elif child.tag=='a':
            href=child.get('href','')
            if href.startswith(('https://','http://')):
                inner=f'<link href="{html.escape(href,quote=True)}" color="#155b89">{inner}</link>'
        elif child.tag=='img': inner=''
        s+=inner+html.escape(child.tail or '')
    return s


def image_path(el):
    imgs=el.xpath('.//img')
    return DOCS / imgs[0].get('src') if imgs else None


def body_elements(tree):
    for el in tree:
        if el.tag=='h1': continue
        txt=text_of(el)
        if 'SAMPLE / DEMONSTRATION' in txt and ('cover' in txt.lower() or 'metadata' in txt.lower()) and el.tag=='p':
            continue
        yield el


def cover_pairs():
    return [
        ('Submitted by',META['student_names']),('Registration number',META['registration_numbers']),
        ('Programme',META['programme']),('Department',META['department']),
        ('Institution',META['institution']),('Supervisor',META['supervisor']),
        ('Academic session',META['academic_session']),('Prepared',META['submission_date']),
    ]


def docx_inline(par, el):
    if el.text: par.add_run(el.text)
    for child in el:
        if child.tag=='img': continue
        run=par.add_run(text_of(child))
        run.bold=child.tag in ('strong','b')
        run.italic=child.tag in ('em','i')
        if child.tag=='code': run.font.name='Courier New';run.font.size=Pt(9)
        if child.tag=='a':
            href=child.get('href','')
            if href.startswith(('https://','http://')):
                par.add_run(f' ({href})').font.size=Pt(9)
        if child.tail:par.add_run(child.tail)


def write_docx(path,title,tree):
    doc=Document()
    sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.8)
    sec.left_margin=sec.right_margin=Inches(.8)
    style=doc.styles['Normal'];style.font.name='Times New Roman';style.font.size=Pt(11)
    style.paragraph_format.space_after=Pt(7)
    style.paragraph_format.line_spacing=1.13
    for name in ['Title','Heading 1','Heading 2','Heading 3']:
        doc.styles[name].font.name='Arial'
        doc.styles[name].font.color.rgb=RGBColor.from_string('16344A')
    footer=sec.footer.paragraphs[0]
    footer.add_run('WebShield AI · Sample academic cover · ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    footer.style=doc.styles['Caption'];footer.alignment=2
    p=doc.add_paragraph('FINAL YEAR PROJECT',style='Subtitle');p.alignment=1
    p=doc.add_paragraph(title,style='Title');p.alignment=1
    p=doc.add_paragraph('SAMPLE / DEMONSTRATION COVER DETAILS');p.alignment=1
    for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor.from_string('AA3829')
    doc.add_paragraph('Dummy identities requested for presentation. These names do not identify an actual student, supervisor or institution.')
    table=doc.add_table(rows=0,cols=2);table.style='Light Shading Accent 1'
    for key,value in cover_pairs():
        cells=table.add_row().cells;cells[0].text=key;cells[1].text=value
    doc.add_paragraph(META['approval_status'])
    doc.add_page_break()
    for el in body_elements(tree):
        tag=el.tag
        if tag in ('h2','h3','h4'):
            if title.endswith('Report') and tag == 'h2' and (re.match(r'^\d+ ', text_of(el)) or text_of(el) == 'Contents'):
                doc.add_page_break()
            doc.add_heading(text_of(el),level=int(tag[1])-1)
        elif tag=='p':
            img=image_path(el)
            if img:
                with Image.open(img) as im:w,h=im.size
                width=min(6.6,5.8*w/h)
                doc.add_picture(str(img),width=Inches(width))
                doc.add_paragraph(el.xpath('.//img')[0].get('alt',''),style='Caption')
            else:docx_inline(doc.add_paragraph(),el)
        elif tag in ('ol','ul'):
            for item in el:
                docx_inline(doc.add_paragraph(style='List Number' if tag=='ol' else 'List Bullet'),item)
        elif tag=='pre':
            p=doc.add_paragraph()
            p.paragraph_format.space_before=Pt(4)
            r=p.add_run(text_of(el));r.font.name='Courier New';r.font.size=Pt(8.5)
        elif tag=='table':
            rows=el.xpath('.//tr');n=max(len(r) for r in rows)
            table=doc.add_table(rows=0,cols=n);table.style='Light Shading Accent 1'
            for row in rows:
                cells=table.add_row().cells
                for i,cell in enumerate(row):
                    docx_inline(cells[i].paragraphs[0],cell)
                    for run in cells[i].paragraphs[0].runs:run.font.size=Pt(9)
                    if cell.tag=='th':
                        for run in cells[i].paragraphs[0].runs:run.bold=True
            doc.add_paragraph()
    doc.core_properties.title=title
    doc.core_properties.subject='Controlled FYP prototype; sample identity details'
    doc.core_properties.author='Sample Student (demonstration metadata)'
    doc.save(path)


class AcademicPDF(SimpleDocTemplate):
    def afterFlowable(self, flowable):
        if isinstance(flowable, Paragraph) and flowable.style.name == 'Heading1':
            label = flowable.getPlainText()
            bookmark = f'heading-{self.page}-{hashlib.sha256(label.encode()).hexdigest()[:10]}'
            self.canv.bookmarkPage(bookmark)
            self.canv.addOutlineEntry(label, bookmark, 0)
            self.notify('TOCEntry', (0, label, self.page, bookmark))


def write_pdf(path,title,tree):
    fonts={'Academic':'Times New Roman.ttf','Academic-Bold':'Times New Roman Bold.ttf','Academic-Italic':'Times New Roman Italic.ttf'}
    active_font_dir = FONT_DIR
    if not all((active_font_dir/file).exists() for file in fonts.values()):
        active_font_dir = Path(reportlab.__file__).parent / 'fonts'
        fonts = {'Academic':'Vera.ttf','Academic-Bold':'VeraBd.ttf','Academic-Italic':'VeraIt.ttf'}
    if all((active_font_dir/file).exists() for file in fonts.values()):
        for name,file in fonts.items():pdfmetrics.registerFont(TTFont(name,str(active_font_dir/file)))
        pdfmetrics.registerFontFamily('Academic',normal='Academic',bold='Academic-Bold',italic='Academic-Italic',boldItalic='Academic-Bold')
        main='Academic'
    else:main='Times-Roman'
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Text',fontName=main,fontSize=11,leading=15,spaceAfter=8))
    styles.add(ParagraphStyle(name='Cell',parent=styles['Text'],fontSize=9,leading=11,spaceAfter=0))
    styles.add(ParagraphStyle(name='CaptionText',fontName=main,fontSize=9,leading=12,textColor=colors.HexColor('#445b6d'),spaceAfter=12))
    styles.add(ParagraphStyle(name='Subtitle',fontName='Helvetica',fontSize=12,leading=16,alignment=TA_CENTER,spaceAfter=20))
    styles.add(ParagraphStyle(name='CoverTitle',fontName='Helvetica-Bold',fontSize=27,leading=34,alignment=TA_CENTER,spaceAfter=24,textColor=colors.HexColor('#16344a')))
    styles.add(ParagraphStyle(name='Sample',fontName='Helvetica-Bold',fontSize=13,leading=18,alignment=TA_CENTER,textColor=colors.HexColor('#aa3829'),spaceAfter=18))
    for name in ('Heading1','Heading2','Heading3'):
        styles[name].textColor=colors.HexColor('#16344a');styles[name].spaceBefore=15;styles[name].spaceAfter=9
    doc=AcademicPDF(str(path),pagesize=A4,rightMargin=20*mm,leftMargin=20*mm,topMargin=19*mm,bottomMargin=20*mm,title=title,author='Sample Student (demonstration metadata)')
    width=A4[0]-40*mm
    P=lambda txt,style='Text':Paragraph(txt,styles[style])
    story=[Spacer(1,35*mm),P('FINAL YEAR PROJECT','Subtitle'),P(html.escape(title),'CoverTitle'),P('SAMPLE / DEMONSTRATION COVER DETAILS','Sample'),P('Dummy identities requested for presentation. These names do not identify an actual student, supervisor or institution.')]
    pairs=[[P(k,'Cell'),P(html.escape(v),'Cell')] for k,v in cover_pairs()]
    table=Table(pairs,colWidths=[width*.32,width*.68],hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(0,-1),colors.HexColor('#edf4f7')),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#c6d2da')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),9),('RIGHTPADDING',(0,0),(-1,-1),9),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
    story += [table,Spacer(1,14*mm),P(html.escape(META['approval_status'])),PageBreak()]
    skip_contents_list = False
    for el in body_elements(tree):
        tag=el.tag
        if tag == 'h2' and text_of(el) == 'Contents':
            toc = TableOfContents()
            toc.levelStyles = [ParagraphStyle(name='ContentsEntry', fontName=main, fontSize=11, leading=20, leftIndent=0, firstLineIndent=0)]
            story.extend([PageBreak(), P('Contents', 'Title'), Spacer(1, 8*mm), toc, PageBreak()])
            skip_contents_list = True
            continue
        if skip_contents_list and tag == 'ol':
            skip_contents_list = False
            continue
        if tag in ('h2','h3','h4'):
            if title.endswith('Report') and tag == 'h2' and re.match(r'^\d+ ', text_of(el)) and not isinstance(story[-1], PageBreak):
                story.append(PageBreak())
            story.append(P(inline(el),f'Heading{int(tag[1])-1}'))
        elif tag=='p':
            img=image_path(el)
            if img:
                with Image.open(img) as im:w,h=im.size
                scale=min(width/w,165*mm/h)
                figure=PDFImage(str(img),width=w*scale,height=h*scale)
                figure.hAlign='CENTER'
                story.append(KeepTogether([figure,Spacer(1,3*mm),P(html.escape(el.xpath('.//img')[0].get('alt','')),'CaptionText')]))
            else:story.append(P(inline(el)))
        elif tag in ('ol','ul'):
            for i,item in enumerate(el):
                story.append(P((f'{i+1}. ' if tag=='ol' else '• ')+inline(item)))
        elif tag=='pre':
            story.append(Preformatted(text_of(el),ParagraphStyle(name='Code',fontName='Courier',fontSize=8,leading=11,spaceAfter=10),maxLineLength=90))
        elif tag=='table':
            rows=el.xpath('.//tr');n=max(len(r) for r in rows)
            matrix=[[P(inline(cell),'Cell') for cell in row] for row in rows]
            tbl=Table(matrix,colWidths=[width/n]*n,repeatRows=1,hAlign='LEFT')
            tbl.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7f0f4')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f7fafb')]),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#c6d2da')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
            story.extend([tbl,Spacer(1,10)])
    def footer(canvas,doc):
        canvas.saveState();canvas.setFillColor(colors.white);canvas.rect(0,0,A4[0],A4[1],fill=1,stroke=0);canvas.setStrokeColor(colors.HexColor('#c6d2da'));canvas.line(20*mm,16*mm,A4[0]-20*mm,16*mm)
        canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#445b6d'))
        canvas.drawString(20*mm,11*mm,'WebShield AI · Sample academic cover')
        canvas.drawRightString(A4[0]-20*mm,11*mm,f'{doc.page}')
        canvas.restoreState()
    doc.multiBuild(story,onFirstPage=footer,onLaterPages=footer)


def write_html(path,title,tree):
    tree=html_parser.fromstring(html_parser.tostring(tree))
    for img in tree.xpath('.//img'):
        asset=DOCS/img.get('src')
        img.set('src','data:image/png;base64,'+base64.b64encode(asset.read_bytes()).decode())
    headings = [el for el in tree if el.tag == 'h2' and re.match(r'^\d+ ', text_of(el))]
    for el in tree:
        if el.tag == 'h2' and text_of(el) == 'Contents':
            listing = el.getnext()
            if listing is not None and listing.tag == 'ol':
                for item, heading in zip(listing, headings):
                    label = text_of(item)
                    item.clear()
                    link = html_parser.Element('a', href='#'+heading.get('id',''))
                    link.text = label
                    item.append(link)
    cover='<section class="cover"><p>FINAL YEAR PROJECT</p><h1>'+html.escape(title)+'</h1><p class="sample">SAMPLE / DEMONSTRATION COVER DETAILS</p><p>Dummy identities requested for presentation. No actual identity or approval is represented.</p><table>'+''.join(f'<tr><th>{html.escape(k)}</th><td>{html.escape(v)}</td></tr>' for k,v in cover_pairs())+'</table><p>'+html.escape(META['approval_status'])+'</p></section>'
    body=''.join(html_parser.tostring(el,encoding='unicode') for el in body_elements(tree))
    style='''body{font:16px/1.6 Georgia,serif;color:#183347;max-width:900px;margin:40px auto;padding:0 28px}h1,h2,h3{font-family:Arial,sans-serif;line-height:1.25;color:#16344a}h2{margin-top:2em}p{margin:0 0 1em}table{border-collapse:collapse;width:100%;font-size:14px;margin:1em 0}td,th{border:1px solid #bdccd5;padding:9px;vertical-align:top;text-align:left;overflow-wrap:anywhere}th{background:#e8f0f4}img{max-width:100%;max-height:850px;object-fit:contain;display:block;margin:1em auto}pre{white-space:pre-wrap;font-size:13px;background:#f1f5f7;padding:14px}a{color:#155b89;overflow-wrap:anywhere}.sample{color:#aa3829;font:bold 18px Arial}.cover{text-align:center;min-height:700px;display:flex;flex-direction:column;justify-content:center;border-bottom:2px solid #c6d2da}.cover table{text-align:left}footer{margin-top:3em;font-size:13px;color:#52697a}@media print{body{font-size:11pt;max-width:none;margin:0}.cover{break-after:page;border:none}h2,h3{break-after:avoid}tr,img{break-inside:avoid}thead{display:table-header-group}}'''
    path.write_text('<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>'+html.escape(title)+'</title><style>'+style+'</style><body>'+cover+'<main>'+body+'</main><footer>WebShield AI · Sample academic cover · No signed academic acceptance</footer></body></html>')


def main():
    OUT.mkdir(exist_ok=True);ASSETS.mkdir(exist_ok=True)
    diagram_assets()
    outputs=[]
    sources=[('FINAL_REPORT.md','WebShield_AI_Final_Report','WebShield AI Final Year Project Report'),('USER_MANUAL.md','WebShield_AI_User_Manual','WebShield AI User Manual')]
    for source,stem,title in sources:
        path=DOCS/source;tree=source_tree(path)
        for ext,fn in [('docx',write_docx),('pdf',write_pdf),('html',write_html)]:
            dest=OUT/f'{stem}.{ext}';fn(dest,title,tree)
            outputs.append({'file':str(dest.relative_to(ROOT)),'bytes':dest.stat().st_size,'sha256':sha(dest),'source':str(path.relative_to(ROOT)),'source_sha256':sha(path)})
    manifest={'identity_status':META['identity_status'],'cover_metadata_sha256':sha(OUT/'sample-cover-metadata.json'),'approval_status':META['approval_status'],'outputs':outputs,'preserved_proposals':[{ 'file':str(p.relative_to(ROOT)),'sha256':sha(p)} for p in sorted(OUT.glob('WebShield_AI_FYP_Proposal.*'))],'assets':[{ 'file':str(p.relative_to(ROOT)),'sha256':sha(p)} for p in sorted(ASSETS.glob('*.png'))]}
    manifest['builder'] = {
        'file': str(Path(__file__).resolve().relative_to(ROOT)),
        'sha256': sha(Path(__file__).resolve()),
        'requirements_file': 'docs/submission/build-requirements.txt',
        'requirements_sha256': sha(OUT/'build-requirements.txt'),
    }
    manifest['companion_documents'] = [
        {'file': str(p.relative_to(ROOT)), 'sha256': sha(p)}
        for p in [DOCS/'DEPLOYMENT_CHECKLIST.md', DOCS/'ACCEPTANCE_CHECKLIST.md', OUT/'README.md']
    ]
    manifest['capture_metadata_sha256'] = sha(ASSETS/'submission-screenshots.json')
    evidence_names = [
        'phreshphish-provenance.json', 'phreshphish-evaluation.json',
        'scamferret-provenance.json', 'scamferret-ablations.json',
        'benchmark-provenance.json', 'benchmark-evaluation.json',
        'model-development-checks.json', 'installed-model-workflow-v2.json',
        'live-acceptance.json', 'model-quality-study.json',
        'operational-acceptance.json', 'aligned-evaluation.json',
        'external-model-validation.json', 'final-research-review.md',
    ]
    manifest['referenced_evidence'] = [
        {'file': str(p.relative_to(ROOT)), 'sha256': sha(p)}
        for name in evidence_names if (p := DOCS/'evidence'/name).exists()
    ]
    (OUT/'submission-manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    for item in outputs:print(item['file'],item['bytes'])


if __name__=='__main__':main()
